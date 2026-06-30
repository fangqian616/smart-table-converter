#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智能表格格式转换系统 - Streamlit Web应用
上传源数据 + 目标模板/指令，AI自动生成映射规则，一键转换输出
"""

import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import json
import os
import tempfile
import time
from datetime import datetime
from pathlib import Path

# ═══════════════════════════════════════
# i18n
# ═══════════════════════════════════════
LANG = {
    "zh": {
        # sidebar
        "sidebar_title": "表格格式转换",
        "sidebar_caption": "上传源数据 + 指定目标格式，AI自动映射转换",
        "ai_section": "AI映射（可选）",
        "ai_key_label": "DeepSeek API Key",
        "ai_key_help": "填入后可AI自动生成映射规则，留空则需手动配置",
        "api_base_label": "API Base URL",
        "api_base_help": "默认DeepSeek官方，也可替换兼容接口",
        "preset_section": "预设管理",
        "load_preset": "加载预设",
        "no_preset": "（不使用）",
        "no_preset_hint": "暂无预设，转换后可保存",
        # main
        "main_title": "🔄 智能表格格式转换系统",
        "main_desc": "上传源数据表格，指定目标格式（模板/指令/预设），AI自动生成列映射规则，预览确认后一键转换输出。",
        # step1
        "step1_title": "① 上传源数据",
        "upload_source": "上传源数据文件",
        "upload_source_help": "支持 Excel 和 CSV 格式",
        "read_success": "读取成功：{} 行 × {} 列",
        "col_list": "列名列表",
        "sample_preview": "前5行预览",
        "read_fail": "读取失败：{}",
        # step2
        "step2_title": "② 指定目标格式",
        "target_mode_label": "目标格式来源",
        "mode_instruction": "📝 自然语言指令",
        "mode_template": "📎 上传模板文件",
        "mode_preset": "💾 使用预设",
        "instruction_placeholder": "例：输出列：订单号(文本)、日期(YYYY-MM-DD)、客户名、金额(保留2位小数)、品类、备注\n日期从YYYYMMDD转成YYYY-MM-DD，金额四舍五入保留2位，按日期升序排列",
        "instruction_help": "用自然语言描述目标列名、格式要求和转换规则",
        "instruction_label": "描述目标格式",
        "template_label": "上传目标模板",
        "template_help": "上传目标格式的Excel模板，系统自动读取列头",
        "template_ok": "模板读取成功",
        "template_cols": "Sheet「{}」列头：",
        "template_fail": "模板读取失败：{}",
        "preset_loaded": "已加载预设「{}」",
        "preset_desc": "描述：{}",
        "preset_none_desc": "无",
        "view_rules": "查看映射规则",
        "preset_select_hint": "请在侧边栏选择一个预设",
        # step3
        "step3_title": "③ 映射规则",
        "ai_need_key": "AI映射需要DeepSeek API Key，请在侧边栏填入。你也可以在下方手动编辑JSON规则。",
        "ai_btn": "🤖 AI生成映射规则",
        "ai_running": "AI正在分析列映射...",
        "ai_ok": "映射规则生成完成！",
        "ai_fail": "AI生成失败：{}",
        "col_mapping_title": "列映射对照",
        "src_col": "源列",
        "target_col": "目标列",
        "transform": "转换",
        "status": "状态",
        "status_ok": "✅ 直接映射",
        "status_warn": "⚠️ 列名不匹配",
        "status_miss": "❌ 未指定源列",
        "status_compute": "🧮 计算生成",
        "empty_src": "（空）",
        "edit_rules": "编辑映射规则（JSON）",
        "rules_json_label": "映射规则JSON",
        "apply_edit": "应用修改",
        "rules_updated": "规则已更新",
        "json_error": "JSON格式错误：{}",
        "manual_hint": "没有AI生成的规则？可以在下方手动输入映射规则JSON：",
        "manual_label": "手动输入映射规则",
        "apply_manual": "应用手动规则",
        "manual_applied": "规则已应用",
        # step4
        "step4_title": "④ 预览与转换",
        "source_rows": "源数据行数",
        "result_rows": "转换后行数",
        "target_cols": "目标列数",
        "preview_title": "转换结果预览（前20行）",
        "gen_excel_btn": "📥 生成Excel文件",
        "gen_excel_running": "正在生成...",
        "dl_excel_btn": "⬇️ 下载转换结果",
        "dl_csv_btn": "⬇️ 下载CSV",
        "gen_csv_btn": "📥 导出CSV文件",
        "output_filename": "转换结果_{}",
        # preset save
        "save_preset_title": "💾 保存为预设",
        "preset_name_label": "预设名称",
        "preset_name_ph": "例：发票格式转换",
        "preset_desc_label": "预设描述",
        "preset_desc_ph": "简要说明用途",
        "save_btn": "保存预设",
        "preset_saved": "预设「{}」已保存！",
        "convert_fail": "转换执行失败：{}",
        "error_detail": "错误详情",
        # intro
        "intro_title": "📤 使用流程",
        "intro_1": "1. **上传源数据** — Excel/CSV文件",
        "intro_2": "2. **指定目标格式** — 三种方式：",
        "intro_2a": "   - 📝 自然语言指令描述目标列和格式要求",
        "intro_2b": "   - 📎 上传目标模板Excel，自动读取列头",
        "intro_2c": "   - 💾 使用之前保存的映射预设",
        "intro_3": "3. **AI生成映射** — 自动匹配源列→目标列+转换规则",
        "intro_4": "4. **预览确认** — 查看映射对照表和转换预览",
        "intro_5": "5. **一键导出** — 下载Excel或CSV",
        "intro_features": "**支持的转换：**",
        "intro_f1": "- 列名映射、格式转换（日期/数字/文本）",
        "intro_f2": "- 新增计算列、行过滤、排序",
        "intro_f3": "- 基于模板保留原格式输出",
    },
    "en": {
        # sidebar
        "sidebar_title": "Table Converter",
        "sidebar_caption": "Upload source data + target format, AI auto-mapping",
        "ai_section": "AI Mapping (Optional)",
        "ai_key_label": "DeepSeek API Key",
        "ai_key_help": "Enter to enable AI auto-mapping; leave empty for manual config",
        "api_base_label": "API Base URL",
        "api_base_help": "Default: DeepSeek official; replace with compatible endpoint",
        "preset_section": "Presets",
        "load_preset": "Load Preset",
        "no_preset": "(None)",
        "no_preset_hint": "No presets yet; save one after conversion",
        # main
        "main_title": "🔄 Smart Table Format Converter",
        "main_desc": "Upload source data, specify target format (template/instructions/preset), AI auto-generates column mapping rules, preview and convert.",
        # step1
        "step1_title": "① Upload Source Data",
        "upload_source": "Upload source data file",
        "upload_source_help": "Supports Excel and CSV formats",
        "read_success": "Read success: {} rows × {} columns",
        "col_list": "Columns",
        "sample_preview": "First 5 rows preview",
        "read_fail": "Read failed: {}",
        # step2
        "step2_title": "② Specify Target Format",
        "target_mode_label": "Target format source",
        "mode_instruction": "📝 Natural Language Instructions",
        "mode_template": "📎 Upload Template",
        "mode_preset": "💾 Use Preset",
        "instruction_placeholder": "e.g. Output columns: Order No. (text), Date (YYYY-MM-DD), Customer, Amount (2 decimals), Category, Notes\nConvert date from YYYYMMDD to YYYY-MM-DD, round amount to 2 decimals, sort by date ascending",
        "instruction_help": "Describe target column names, format requirements, and transform rules in plain text",
        "instruction_label": "Describe target format",
        "template_label": "Upload target template",
        "template_help": "Upload target Excel template; column headers will be auto-read",
        "template_ok": "Template read successfully",
        "template_cols": "Sheet '{}' columns:",
        "template_fail": "Template read failed: {}",
        "preset_loaded": "Loaded preset '{}'",
        "preset_desc": "Description: {}",
        "preset_none_desc": "None",
        "view_rules": "View mapping rules",
        "preset_select_hint": "Please select a preset from the sidebar",
        # step3
        "step3_title": "③ Mapping Rules",
        "ai_need_key": "AI mapping requires a DeepSeek API Key in the sidebar. You can also manually edit JSON rules below.",
        "ai_btn": "🤖 AI Generate Mapping Rules",
        "ai_running": "AI is analyzing column mapping...",
        "ai_ok": "Mapping rules generated!",
        "ai_fail": "AI generation failed: {}",
        "col_mapping_title": "Column Mapping",
        "src_col": "Source",
        "target_col": "Target",
        "transform": "Transform",
        "status": "Status",
        "status_ok": "✅ Direct mapping",
        "status_warn": "⚠️ Column mismatch",
        "status_miss": "❌ No source column",
        "status_compute": "🧮 Computed",
        "empty_src": "(empty)",
        "edit_rules": "Edit mapping rules (JSON)",
        "rules_json_label": "Mapping rules JSON",
        "apply_edit": "Apply changes",
        "rules_updated": "Rules updated",
        "json_error": "JSON format error: {}",
        "manual_hint": "No AI-generated rules? Enter mapping rules JSON manually below:",
        "manual_label": "Enter mapping rules manually",
        "apply_manual": "Apply manual rules",
        "manual_applied": "Rules applied",
        # step4
        "step4_title": "④ Preview & Convert",
        "source_rows": "Source rows",
        "result_rows": "Result rows",
        "target_cols": "Target columns",
        "preview_title": "Conversion preview (first 20 rows)",
        "gen_excel_btn": "📥 Generate Excel",
        "gen_excel_running": "Generating...",
        "dl_excel_btn": "⬇️ Download Excel",
        "dl_csv_btn": "⬇️ Download CSV",
        "gen_csv_btn": "📥 Export CSV",
        "output_filename": "converted_{}",
        # preset save
        "save_preset_title": "💾 Save as Preset",
        "preset_name_label": "Preset name",
        "preset_name_ph": "e.g. Invoice Format",
        "preset_desc_label": "Preset description",
        "preset_desc_ph": "Brief description of usage",
        "save_btn": "Save Preset",
        "preset_saved": "Preset '{}' saved!",
        "convert_fail": "Conversion failed: {}",
        "error_detail": "Error details",
        # intro
        "intro_title": "📤 How to Use",
        "intro_1": "1. **Upload source data** — Excel or CSV file",
        "intro_2": "2. **Specify target format** — 3 options:",
        "intro_2a": "   - 📝 Natural language instructions for target columns and format",
        "intro_2b": "   - 📎 Upload target template to auto-read column headers",
        "intro_2c": "   - 💾 Load previously saved mapping preset",
        "intro_3": "3. **AI generates mapping** — Auto-match source → target columns + transform rules",
        "intro_4": "4. **Preview & confirm** — View mapping table and converted data",
        "intro_5": "5. **Export** — Download Excel or CSV",
        "intro_features": "**Supported transforms:**",
        "intro_f1": "- Column mapping, format conversion (date/number/text)",
        "intro_f2": "- Computed columns, row filtering, sorting",
        "intro_f3": "- Template-based output preserving original formatting",
    },
}

def t(key, *args):
    """Get translated text by key"""
    lang = st.session_state.get("lang", "zh")
    text = LANG.get(lang, LANG["zh"]).get(key, key)
    if args:
        text = text.format(*args)
    return text

# ── 页面配置 ──
st.set_page_config(
    page_title="Smart Table Converter",
    page_icon="🔄",
    layout="wide",
)

# ── 常量 ──
PRESET_DIR = Path(__file__).parent / "presets"
PRESET_DIR.mkdir(exist_ok=True)
SAMPLE_ROWS = 5

# ── 自定义CSS ──
st.markdown("""
<style>
    .main-title { font-size: 2rem; font-weight: 700; color: #2C3E50; }
    .mapping-ok { color: #27AE60; font-weight: 600; }
    .mapping-warn { color: #E67E22; font-weight: 600; }
    .mapping-miss { color: #E74C3C; font-weight: 600; }
</style>
""", unsafe_allow_html=True)

# ═══════════════════════════════════════
# 工具函数
# ═══════════════════════════════════════

def read_source_file(file_obj, encoding="utf-8"):
    name = file_obj.name.lower()
    suffix = Path(name).suffix
    if suffix == ".csv":
        return pd.read_csv(file_obj, encoding=encoding)
    elif suffix in (".xls", ".xlsx", ".xlsm"):
        return pd.read_excel(file_obj, engine="openpyxl" if suffix != ".xls" else None)
    else:
        try:
            return pd.read_excel(file_obj)
        except:
            return pd.read_csv(file_obj, encoding=encoding)


def read_template_columns(file_obj):
    wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    result = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            if cell.value is not None:
                headers.append({"column": cell.column, "name": str(cell.value)})
        if headers:
            result[sheet_name] = headers
    wb.close()
    return result


def get_sample_data(df, n=SAMPLE_ROWS):
    rows = []
    for _, row in df.head(n).iterrows():
        row_dict = {}
        for col in df.columns:
            val = row[col]
            row_dict[str(col)] = str(val) if pd.notna(val) else ""
        rows.append(row_dict)
    return rows


def call_deepseek_for_mapping(api_key, api_base, source_columns, source_sample,
                               target_columns=None, instruction=None):
    import requests
    
    lang = st.session_state.get("lang", "zh")
    
    source_info = {"columns": source_columns, "sample_data": source_sample}
    
    target_info = ""
    if target_columns:
        target_info = f"\nTarget columns:\n{json.dumps(target_columns, ensure_ascii=False, indent=2)}"
    
    instruction_text = ""
    if instruction and instruction.strip():
        instruction_text = f"\nUser instructions:\n{instruction.strip()}"
    
    if lang == "en":
        prompt = f"""You are a data format conversion expert. Generate column mapping rules based on source data columns and target format.

Source data info:
Column names: {json.dumps(source_columns, ensure_ascii=False)}
Sample data (first {len(source_sample)} rows):
{json.dumps(source_sample, ensure_ascii=False, indent=2)}
{target_info}{instruction_text}

Output strict JSON mapping rules only, no other text. Format:
{{
  "mappings": [
    {{"source": "source_column", "target": "target_column", "transform": "transform_type"}}
  ],
  "computed": [
    {{"target": "target_column", "formula": "formula_description"}}
  ],
  "filters": [],
  "sort": {{"column": "sort_column", "order": "asc or desc"}}
}}

Available transforms:
- "none": copy as-is
- "date:SRC->DST": date format conversion, e.g. "date:YYYYMMDD->YYYY-MM-DD"
- "round:N": round to N decimal places
- "text": convert to text
- "strip": trim whitespace
- "upper"/"lower": uppercase/lowercase
- "prefix:X"/"suffix:X": add prefix/suffix
- "regex:PATTERN:REPLACEMENT": regex replacement

If a target column has no matching source, put it in computed with a formula, or leave source empty.
Output JSON only, no explanation."""
    else:
        prompt = f"""你是数据格式转换专家。根据源数据列和目标格式，生成列映射规则。

源数据信息:
列名列表: {json.dumps(source_columns, ensure_ascii=False)}
样本数据(前{len(source_sample)}行):
{json.dumps(source_sample, ensure_ascii=False, indent=2)}
{target_info}{instruction_text}

请输出严格的JSON格式映射规则，不要输出任何其他内容。格式如下:
{{
  "mappings": [
    {{"source": "源列名", "target": "目标列名", "transform": "转换方式"}}
  ],
  "computed": [
    {{"target": "目标列名", "formula": "计算公式描述"}}
  ],
  "filters": [],
  "sort": {{"column": "排序列", "order": "asc或desc"}}
}}

transform可选值:
- "none": 直接复制
- "date:源格式->目标格式": 日期格式转换
- "round:N": 四舍五入保留N位小数
- "text": 转为文本格式
- "strip": 去除首尾空格
- "upper"/"lower": 大写/小写
- "prefix:前缀"/"suffix:后缀": 添加前缀/后缀
- "regex:正则表达式:替换值": 正则替换

如果某目标列在源数据中找不到对应，放入computed用formula描述如何计算，或留空source。
只输出JSON，不要输出解释。"""

    resp = requests.post(
        f"{api_base}/chat/completions",
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json={
            "model": "deepseek-chat",
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.3,
            "max_tokens": 2000,
        },
        timeout=60,
    )
    
    if resp.status_code != 200:
        raise Exception(f"API call failed: {resp.status_code} - {resp.text}")
    
    content = resp.json()['choices'][0]['message']['content'].strip()
    if content.startswith("```"):
        content = content.split("\n", 1)[1] if "\n" in content else content[3:]
        if content.endswith("```"):
            content = content[:-3]
    content = content.strip()
    return json.loads(content)


def apply_transform(value, transform):
    if transform == "none" or not transform:
        return value
    if transform == "text":
        return str(value) if value is not None else ""
    if transform == "strip":
        return str(value).strip() if value is not None else ""
    if transform == "upper":
        return str(value).upper() if value is not None else ""
    if transform == "lower":
        return str(value).lower() if value is not None else ""
    if transform.startswith("round:"):
        n = int(transform.split(":")[1])
        try: return round(float(value), n)
        except: return value
    if transform.startswith("date:"):
        fmt_part = transform[5:]
        if "->" in fmt_part:
            src_fmt, dst_fmt = fmt_part.split("->", 1)
            src_fmt = src_fmt.replace("YYYY", "%Y").replace("MM", "%m").replace("DD", "%d")
            dst_fmt = dst_fmt.replace("YYYY", "%Y").replace("MM", "%m").replace("DD", "%d")
            try:
                from datetime import datetime as dt
                if isinstance(value, (pd.Timestamp,)):
                    return value.strftime(dst_fmt)
                return dt.strptime(str(value), src_fmt).strftime(dst_fmt)
            except: return value
    if transform.startswith("prefix:"):
        prefix = transform.split(":", 1)[1]
        return f"{prefix}{value}" if value is not None else value
    if transform.startswith("suffix:"):
        suffix = transform.split(":", 1)[1]
        return f"{value}{suffix}" if value is not None else value
    if transform.startswith("regex:"):
        parts = transform.split(":", 2)
        if len(parts) >= 3:
            import re
            try: return re.sub(parts[1], parts[2], str(value))
            except: return value
    return value


def execute_conversion(df, mapping_rules):
    mappings = mapping_rules.get("mappings", [])
    computed = mapping_rules.get("computed", [])
    filters = mapping_rules.get("filters", [])
    sort_config = mapping_rules.get("sort", {})
    result_data = {}
    target_columns = []
    for m in mappings:
        target = m.get("target", "")
        source = m.get("source", "")
        transform = m.get("transform", "none")
        if not target: continue
        target_columns.append(target)
        if source and source in df.columns:
            result_data[target] = df[source].apply(lambda v: apply_transform(v, transform))
        else:
            result_data[target] = ""
    for c in computed:
        target = c.get("target", "")
        formula = c.get("formula", "")
        if not target: continue
        if target not in target_columns: target_columns.append(target)
        result_data[target] = _eval_formula(formula, df, result_data)
    result_df = pd.DataFrame(result_data, columns=target_columns)
    for f in filters:
        col = f.get("column", "")
        op = f.get("op", "eq")
        val = f.get("value", "")
        if col and col in result_df.columns:
            if op == "eq": result_df = result_df[result_df[col].astype(str) == str(val)]
            elif op == "neq": result_df = result_df[result_df[col].astype(str) != str(val)]
            elif op == "contains": result_df = result_df[result_df[col].astype(str).str.contains(str(val), na=False)]
            elif op == "not_empty": result_df = result_df[result_df[col].astype(str) != ""]
    if sort_config:
        sort_col = sort_config.get("column", "")
        sort_order = sort_config.get("order", "asc")
        if sort_col and sort_col in result_df.columns:
            result_df = result_df.sort_values(by=sort_col, ascending=(sort_order == "asc"))
    return result_df.reset_index(drop=True)


def _eval_formula(formula, source_df, result_data):
    import re
    if formula.startswith("concat(") and formula.endswith(")"):
        inner = formula[7:-1]
        parts = [p.strip().strip("'\"") for p in inner.split(",")]
        series_list = []
        for p in parts:
            if p in source_df.columns:
                series_list.append(source_df[p].astype(str))
            elif p in result_data:
                series_list.append(result_data[p].astype(str) if isinstance(result_data[p], pd.Series) else pd.Series([str(result_data[p])]*len(source_df)))
            else:
                series_list.append(pd.Series([p]*len(source_df)))
        result = series_list[0]
        for s in series_list[1:]:
            result = result + s
        return result
    if formula in source_df.columns:
        return source_df[formula]
    if formula in result_data and isinstance(result_data[formula], pd.Series):
        return result_data[formula]
    return pd.Series([formula]*len(source_df))


def save_preset(name, mapping_rules, description=""):
    preset = {"name": name, "description": description, "rules": mapping_rules, "created_at": datetime.now().strftime("%Y-%m-%d %H:%M")}
    filepath = PRESET_DIR / f"{name}.json"
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(preset, f, ensure_ascii=False, indent=2)
    return filepath


def load_presets():
    presets = []
    for fp in PRESET_DIR.glob("*.json"):
        with open(fp, "r", encoding="utf-8") as f:
            presets.append(json.load(f))
    return presets


def delete_preset(name):
    filepath = PRESET_DIR / f"{name}.json"
    if filepath.exists():
        filepath.unlink()


def generate_output_excel(df, template_path=None):
    tmp = tempfile.mktemp(suffix=".xlsx")
    if template_path:
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        for r_idx, row in df.iterrows():
            for c_idx, val in enumerate(row):
                ws.cell(row=r_idx + 2, column=c_idx + 1, value=val)
        wb.save(tmp)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = t("result_sheet", "Converted")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        for c_idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=c_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        for r_idx, row in df.iterrows():
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx + 2, column=c_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
        for c_idx, col_name in enumerate(df.columns, 1):
            max_len = max(len(str(col_name)), df[col_name].astype(str).str.len().max() if len(df) > 0 else 0)
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 4, 40)
        wb.save(tmp)
    return tmp


# ═══════════════════════════════════════
# 侧边栏
# ═══════════════════════════════════════
with st.sidebar:
    # 语言切换（最顶部）
    lang_col1, lang_col2 = st.columns(2)
    with lang_col1:
        if st.button("🇨🇳 中文", use_container_width=True, disabled=st.session_state.get("lang","zh")=="zh"):
            st.session_state["lang"] = "zh"
            st.rerun()
    with lang_col2:
        if st.button("🇬🇧 EN", use_container_width=True, disabled=st.session_state.get("lang","zh")=="en"):
            st.session_state["lang"] = "en"
            st.rerun()
    
    st.divider()
    
    st.image("https://img.icons8.com/fluency/96/convert-text.png", width=60)
    st.title(t("sidebar_title"))
    st.caption(t("sidebar_caption"))
    
    st.divider()
    
    st.subheader(t("ai_section"))
    deepseek_key = st.text_input(t("ai_key_label"), type="password", help=t("ai_key_help"))
    deepseek_base = st.text_input(t("api_base_label"), value="https://api.deepseek.com/v1", help=t("api_base_help"))
    
    st.divider()
    
    st.subheader(t("preset_section"))
    existing_presets = load_presets()
    preset_names = [p["name"] for p in existing_presets]
    if preset_names:
        selected_preset = st.selectbox(t("load_preset"), [t("no_preset")] + preset_names)
    else:
        selected_preset = t("no_preset")
        st.caption(t("no_preset_hint"))

# ═══════════════════════════════════════
# 主区域
# ═══════════════════════════════════════
st.markdown(f'<p class="main-title">{t("main_title")}</p>', unsafe_allow_html=True)
st.markdown(t("main_desc"))

# ── Step 1 ──
st.header(t("step1_title"))
source_file = st.file_uploader(t("upload_source"), type=["xlsx", "xls", "csv"], help=t("upload_source_help"))

if source_file:
    try:
        source_df = read_source_file(source_file)
        st.success(t("read_success", len(source_df), len(source_df.columns)))
        col1, col2 = st.columns(2)
        with col1:
            st.caption(t("col_list"))
            st.write(list(source_df.columns))
        with col2:
            st.caption(t("sample_preview"))
            st.dataframe(source_df.head(), use_container_width=True, hide_index=True)
    except Exception as e:
        st.error(t("read_fail", e))
        source_df = None
else:
    source_df = None

# ── Step 2 ──
if source_df is not None:
    st.divider()
    st.header(t("step2_title"))
    
    target_mode = st.radio(t("target_mode_label"),
                           [t("mode_instruction"), t("mode_template"), t("mode_preset")],
                           horizontal=True)
    
    target_columns_info = None
    instruction_text = None
    template_file_path = None
    mapping_rules = None
    
    if target_mode == t("mode_instruction"):
        instruction_text = st.text_area(t("instruction_label"), value="", height=120,
                                         placeholder=t("instruction_placeholder"), help=t("instruction_help"))
    
    elif target_mode == t("mode_template"):
        template_file = st.file_uploader(t("template_label"), type=["xlsx", "xls"], key="template_upload", help=t("template_help"))
        if template_file:
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                    tmp.write(template_file.getvalue())
                    template_file_path = tmp.name
                target_columns_info = read_template_columns(template_file)
                st.success(t("template_ok"))
                for sheet, cols in target_columns_info.items():
                    st.caption(t("template_cols", sheet))
                    st.write([c["name"] for c in cols])
            except Exception as e:
                st.error(t("template_fail", e))
    
    elif target_mode == t("mode_preset"):
        if selected_preset != t("no_preset"):
            preset_data = next((p for p in existing_presets if p["name"] == selected_preset), None)
            if preset_data:
                mapping_rules = preset_data["rules"]
                st.success(t("preset_loaded", selected_preset))
                st.caption(t("preset_desc", preset_data.get("description", t("preset_none_desc"))))
                with st.expander(t("view_rules")):
                    st.json(mapping_rules)
        else:
            st.info(t("preset_select_hint"))
    
    # ── Step 3 ──
    st.divider()
    st.header(t("step3_title"))
    
    need_ai = (target_mode in [t("mode_instruction"), t("mode_template")]) and not mapping_rules
    
    if need_ai:
        if not deepseek_key:
            st.warning(t("ai_need_key"))
        
        can_ai = deepseek_key and not (target_mode == t("mode_instruction") and not (instruction_text and instruction_text.strip()))
        if st.button(t("ai_btn"), disabled=not can_ai):
            with st.spinner(t("ai_running")):
                try:
                    source_columns = [str(c) for c in source_df.columns]
                    source_sample = get_sample_data(source_df)
                    mapping_rules = call_deepseek_for_mapping(deepseek_key, deepseek_base, source_columns, source_sample,
                                                               target_columns=target_columns_info, instruction=instruction_text)
                    st.success(t("ai_ok"))
                except Exception as e:
                    st.error(t("ai_fail", e))
                    mapping_rules = None
    
    if mapping_rules:
        st.subheader(t("col_mapping_title"))
        mappings = mapping_rules.get("mappings", [])
        computed = mapping_rules.get("computed", [])
        
        map_data = []
        for m in mappings:
            source = m.get("source", "")
            target = m.get("target", "")
            transform = m.get("transform", "none")
            if source and source in source_df.columns:
                status = t("status_ok")
            elif source:
                status = t("status_warn")
            else:
                status = t("status_miss")
            map_data.append({t("src_col"): source if source else t("empty_src"), "→": "→",
                             t("target_col"): target, t("transform"): transform, t("status"): status})
        for c in computed:
            map_data.append({t("src_col"): t("status_compute").split(" ")[0] if " " in t("status_compute") else "🧮",
                             "→": "→", t("target_col"): c.get("target", ""),
                             t("transform"): c.get("formula", ""), t("status"): t("status_compute")})
        
        if map_data:
            st.dataframe(map_data, use_container_width=True, hide_index=True)
        
        with st.expander(t("edit_rules")):
            rules_json = st.text_area(t("rules_json_label"), value=json.dumps(mapping_rules, ensure_ascii=False, indent=2), height=300, key="rules_json_edit")
            if st.button(t("apply_edit")):
                try:
                    mapping_rules = json.loads(rules_json)
                    st.success(t("rules_updated"))
                    st.rerun()
                except json.JSONDecodeError as e:
                    st.error(t("json_error", e))
    else:
        st.info(t("manual_hint"))
        manual_json = st.text_area(t("manual_label"), value='{\n  "mappings": [],\n  "computed": [],\n  "filters": [],\n  "sort": {}\n}', height=200, key="manual_json")
        if st.button(t("apply_manual")):
            try:
                mapping_rules = json.loads(manual_json)
                st.success(t("manual_applied"))
            except json.JSONDecodeError as e:
                st.error(t("json_error", e))
    
    # ── Step 4 ──
    if mapping_rules:
        st.divider()
        st.header(t("step4_title"))
        
        try:
            result_df = execute_conversion(source_df, mapping_rules)
            
            col1, col2, col3 = st.columns(3)
            with col1: st.metric(t("source_rows"), f"{len(source_df)}")
            with col2: st.metric(t("result_rows"), f"{len(result_df)}")
            with col3: st.metric(t("target_cols"), f"{len(result_df.columns)}")
            
            st.subheader(t("preview_title"))
            st.dataframe(result_df.head(20), use_container_width=True, hide_index=True)
            
            col_dl1, col_dl2 = st.columns(2)
            with col_dl1:
                if st.button(t("gen_excel_btn")):
                    with st.spinner(t("gen_excel_running")):
                        output_path = generate_output_excel(result_df, template_path=template_file_path)
                        with open(output_path, "rb") as f:
                            st.download_button(label=t("dl_excel_btn"), data=f.read(),
                                               file_name=f"{t('output_filename', datetime.now().strftime('%Y%m%d_%H%M'))}.xlsx",
                                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            
            with col_dl2:
                if st.button(t("gen_csv_btn")):
                    csv_data = result_df.to_csv(index=False).encode("utf-8-sig")
                    st.download_button(label=t("dl_csv_btn"), data=csv_data,
                                       file_name=f"{t('output_filename', datetime.now().strftime('%Y%m%d_%H%M'))}.csv", mime="text/csv")
            
            st.divider()
            st.subheader(t("save_preset_title"))
            col_s1, col_s2 = st.columns([1, 2])
            with col_s1: preset_name = st.text_input(t("preset_name_label"), placeholder=t("preset_name_ph"))
            with col_s2: preset_desc = st.text_input(t("preset_desc_label"), placeholder=t("preset_desc_ph"))
            
            if st.button(t("save_btn")) and preset_name:
                save_preset(preset_name, mapping_rules, preset_desc)
                st.success(t("preset_saved", preset_name))
        
        except Exception as e:
            st.error(t("convert_fail", e))
            with st.expander(t("error_detail")):
                st.code(str(e))

else:
    st.divider()
    st.info(f"""### {t('intro_title')}

{t('intro_1')}
{t('intro_2')}
{t('intro_2a')}
{t('intro_2b')}
{t('intro_2c')}
{t('intro_3')}
{t('intro_4')}
{t('intro_5')}

{t('intro_features')}
{t('intro_f1')}
{t('intro_f2')}
{t('intro_f3')}
""")
